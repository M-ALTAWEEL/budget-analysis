"""A bank statement as a list of transactions (feature F11).

Statements come as Excel, CSV or PDF, and rarely as a clean table: a PDF turned into
Excel splits one transaction over several rows, moves the amount into the description, or
lists three dates before their three amounts. So the reader works on the statement as
lines of text, like a person reading it:

1. A **date** starts a transaction (with the bank's reference code next to it).
2. Money is written the Turkish way ("1.234,56", "-875,00") or sits in number cells; on a
   line, the last amount is the **balance** after the transaction and the one before it
   the amount. A line with amounts but no date belongs to the oldest transaction still
   waiting for its amounts; a line with only words continues the last description.
3. Where balances are known, **amount = this balance - the previous one**: exact, and it
   fixes amounts or signs that slid into the wrong place. The carried-over balance at the
   top ("رصيد الترحيل", "Devreden bakiye") starts the chain.

A clean table with headers (date, description, amount, or debit and credit) in any of the
three languages is read directly instead. Rules: docs/BUSINESS_LOGIC.md, section 6.9.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation

MAX_ROWS = 2000

_DATE = re.compile(
    r"(?<!\d)(\d{2})[./](\d{2})[./](\d{4})(?!\d)|(?<!\d)(\d{4})-(\d{2})-(\d{2})(?!\d)"
)
# "1.234,56", "-875,00", "- 12.349,85" (Turkish), never part of a longer number or card.
_TR_MONEY = re.compile(r"(?<![\d.,*])(-\s*)?(\d{1,3}(?:\.\d{3})*|\d+),(\d{2})(?![\d*,])")
_REF = re.compile(r"^[A-Z0-9]{5}$")
_OPENING = ("رصيد الترحيل", "devreden bakiye", "önceki bakiye", "onceki bakiye", "opening balance")
_TOTAL = ("الإجمالي", "الاجمالي", "genel toplam", "toplam", "total", "closing balance")


@dataclass
class Transaction:
    day: date
    description: str
    amount: Decimal  # negative = money out
    balance: Decimal | None = None


@dataclass
class _Open:
    day: date
    words: list[str] = field(default_factory=list)
    numbers: list[Decimal] = field(default_factory=list)


def _tr_money(text: str) -> list[Decimal]:
    found = []
    for sign, whole, cents in _TR_MONEY.findall(text):
        value = Decimal(whole.replace(".", "") + "." + cents)
        found.append(-value if sign else value)
    return found


def _strip_money(text: str) -> str:
    return _TR_MONEY.sub(" ", text)


def _first_date(text: str) -> tuple[date | None, str]:
    match = _DATE.search(text)
    if not match:
        return None, text
    try:
        if match.group(1):
            found = date(int(match.group(3)), int(match.group(2)), int(match.group(1)))
        else:
            found = date(int(match.group(4)), int(match.group(5)), int(match.group(6)))
    except ValueError:
        return None, text
    # Only a date at the start of a line opens a transaction ("8/7/2026" inside a card
    # line is the card's own record, not a new transaction).
    if text[: match.start()].strip():
        return None, text
    return found, text[match.end() :]


def _line(cells: list) -> tuple[date | None, str, list[Decimal]]:
    """One row as (date, words, amounts)."""
    day: date | None = None
    words: list[str] = []
    numbers: list[Decimal] = []
    for i, cell in enumerate(cells):
        if cell is None or cell == "":
            continue
        if isinstance(cell, datetime | date):
            if day is None and i == 0:
                day = cell.date() if isinstance(cell, datetime) else cell
            continue
        if isinstance(cell, int | float | Decimal) and not isinstance(cell, bool):
            numbers.append(Decimal(str(cell)).quantize(Decimal("0.01")))
            continue
        text = str(cell).replace("\n", " ").strip()
        if day is None and not words and not numbers:
            found, rest = _first_date(text)
            if found:
                day, text = found, rest
        numbers.extend(_tr_money(text))
        text = " ".join(_strip_money(text).split())
        if text:
            words.append(text)
    return day, " ".join(words), numbers


def _flow(rows: list[list]) -> list[Transaction]:
    opening: Decimal | None = None
    done: list[_Open] = []
    waiting: list[_Open] = []  # dates whose amounts haven't come yet, oldest first
    for cells in rows:
        day, text, numbers = _line(cells)
        lower = text.lower()
        if day is None and any(word in lower for word in _OPENING):
            opening = numbers[-1] if numbers else opening
            continue
        if day is None and any(word in lower for word in _TOTAL) and len(text) < 40:
            continue  # the closing total line
        if day is not None:
            ref, _, rest = text.partition(" ")
            if _REF.match(ref):
                text = rest.strip()
            item = _Open(day, [text] if text else [], numbers)
            done.append(item)
            if not numbers:
                waiting.append(item)
            continue
        if numbers and waiting:
            item = waiting.pop(0)
            item.numbers = numbers
            if text:
                item.words.append(text)
        elif numbers and done and len(done[-1].numbers) < 2:
            # The amount or the balance went to the next line.
            done[-1].numbers.extend(numbers)
            if text:
                done[-1].words.append(text)
        elif text and done:
            done[-1].words.append(text)

    result: list[Transaction] = []
    previous = opening
    with_numbers = [item for item in done if item.numbers]  # a date alone is dropped
    for i, item in enumerate(with_numbers):
        following = with_numbers[i + 1].numbers if i + 1 < len(with_numbers) else None
        amount, balance = _pair(item.numbers, previous, following)
        if amount == 0:
            continue
        if balance is not None:
            previous = balance
        elif previous is not None:
            previous += amount
        result.append(Transaction(item.day, " ".join(item.words).strip(), amount, balance))
    result.sort(key=lambda t: t.day)  # stable: same-day order is kept
    return result[:MAX_ROWS]


def _pairs(numbers: list[Decimal], previous: Decimal | None) -> list[tuple[Decimal, Decimal]]:
    """Every (amount, balance) where previous + amount = balance, most likely first: the
    amount written before the balance and with its own sign, then any other order and sign
    (a sign that slid onto the balance belongs to the amount)."""
    if previous is None or len(numbers) < 2:
        return []
    tail = numbers[-3:]
    ranked = []
    for b, written in enumerate(tail):
        balance = abs(written)
        for a, amount in enumerate(tail):
            if a == b:
                continue
            for flipped, signed in ((False, amount), (True, -amount)):
                if previous + signed == balance:
                    rank = (a > b) + flipped + (written < 0)
                    ranked.append((rank, signed, balance))
    ranked.sort(key=lambda r: r[0])
    found: list[tuple[Decimal, Decimal]] = []
    for _, signed, balance in ranked:
        if (signed, balance) not in found:
            found.append((signed, balance))
    return found


def _pair(
    numbers: list[Decimal], previous: Decimal | None, following: list[Decimal] | None = None
) -> tuple[Decimal, Decimal | None]:
    """Which number is the amount and which the balance. With the balance before known,
    the pair where previous + amount = balance wins, whatever order they came in. When
    several pairs fit, the one the next transaction's numbers follow from wins ("12.349,85"
    and "12.250,00" after 24.599,85: only one of them leads on); otherwise the most likely.
    With nothing to check against, the last number is the balance and the one before it the
    amount; a single number is the amount."""
    if len(numbers) == 1:
        return numbers[0], None
    fits = _pairs(numbers, previous)
    if len(fits) > 1 and following:
        leading = [fit for fit in fits if _pairs(following, fit[1])]
        if len(leading) == 1:
            return leading[0]
    if fits:
        return fits[0]
    return numbers[-2], numbers[-1]


# ── clean tables ─────────────────────────────────────────────────────────────

_HEADERS = {
    "date": ("tarih", "date", "التاريخ", "işlem tarihi", "islem tarihi", "valör"),
    "description": ("açıklama", "aciklama", "description", "details", "الوصف", "البيان", "işlem"),
    "amount": ("tutar", "amount", "المبلغ", "işlem tutarı"),
    "debit": ("borç", "borc", "debit", "مدين", "withdrawal"),
    "credit": ("alacak", "credit", "دائن", "deposit"),
    "balance": ("bakiye", "balance", "الرصيد"),
}


def _number(value) -> Decimal | None:
    if value is None or value == "":
        return None
    if isinstance(value, int | float | Decimal) and not isinstance(value, bool):
        return Decimal(str(value)).quantize(Decimal("0.01"))
    text = str(value).strip().replace(" ", "").replace("TL", "").replace("₺", "")
    if re.fullmatch(r"-?\d{1,3}(\.\d{3})*,\d{1,2}|-?\d+,\d{1,2}", text):
        text = text.replace(".", "").replace(",", ".")
    else:
        text = text.replace(",", "")
    try:
        return Decimal(text).quantize(Decimal("0.01"))
    except InvalidOperation:
        return None


def _day(value) -> date | None:
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    if value is None:
        return None
    found, _ = _first_date(str(value).strip())
    return found


def _table(rows: list[list]) -> list[Transaction] | None:
    """Headers in the first 15 rows: read the columns they name."""
    for at, row in enumerate(rows[:15]):
        names = [str(c).strip().lower() if c is not None else "" for c in row]
        columns: dict[str, int] = {}
        for key, words in _HEADERS.items():
            for i, name in enumerate(names):
                if name and any(name == w or name.startswith(w) for w in words):
                    columns.setdefault(key, i)
                    break
        if (
            "date" in columns
            and "description" in columns
            and ("amount" in columns or "debit" in columns or "credit" in columns)
        ):
            found: list[Transaction] = []
            for cells in rows[at + 1 :]:
                cells = list(cells) + [None] * (len(names) - len(cells))
                day = _day(cells[columns["date"]])
                if day is None:
                    continue
                if "amount" in columns:
                    amount = _number(cells[columns["amount"]])
                else:
                    out = _number(cells[columns["debit"]]) if "debit" in columns else None
                    into = _number(cells[columns["credit"]]) if "credit" in columns else None
                    amount = (into or Decimal(0)) - abs(out or Decimal(0))
                if not amount:
                    continue
                text = str(cells[columns["description"]] or "").replace("\n", " ").strip()
                balance = _number(cells[columns["balance"]]) if "balance" in columns else None
                found.append(Transaction(day, " ".join(text.split()), amount, balance))
            return found[:MAX_ROWS]
    return None


# ── files ────────────────────────────────────────────────────────────────────


def _rows_from_xlsx(content: bytes) -> list[list]:
    from openpyxl import load_workbook

    book = load_workbook(io.BytesIO(content), read_only=True, data_only=True)
    rows: list[list] = []
    for sheet in book.worksheets:  # in order: a long statement spans several sheets
        rows.extend(list(row) for row in sheet.iter_rows(values_only=True))
    return rows


def _rows_from_csv(content: bytes) -> list[list]:
    for encoding in ("utf-8-sig", "cp1254", "latin-1"):
        try:
            text = content.decode(encoding)
            break
        except UnicodeDecodeError:
            continue
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t") if text.strip() else "excel"
    return [list(row) for row in csv.reader(io.StringIO(text), dialect)]


def _rows_from_pdf(content: bytes) -> list[list]:
    from pypdf import PdfReader

    rows: list[list] = []
    for page in PdfReader(io.BytesIO(content)).pages:
        rows.extend([line] for line in (page.extract_text() or "").splitlines())
    return rows


def read(content: bytes, filename: str) -> list[Transaction]:
    """Raises ValueError when the file can't be read or has no transactions."""
    name = filename.lower()
    try:
        if name.endswith((".xlsx", ".xlsm")):
            rows = _rows_from_xlsx(content)
        elif name.endswith((".csv", ".txt")):
            rows = _rows_from_csv(content)
        elif name.endswith(".pdf"):
            rows = _rows_from_pdf(content)
        else:
            raise ValueError("Use an Excel, CSV or PDF file.")
    except ValueError:
        raise
    except Exception as exc:  # a broken or password-protected file
        raise ValueError("The file couldn't be read.") from exc
    found = _table(rows) or _flow(rows)
    if not found:
        raise ValueError("No transactions were found in the file.")
    return found


# ── what each line is ────────────────────────────────────────────────────────

_CITIES = {"istanbul", "ankara", "izmir", "kocaeli", "bursa", "antalya", "bakırköy", "bakirkoy",
           "konya", "adana", "gaziantep", "eskişehir", "eskisehir", "sakarya"}  # fmt: skip
_LONG_DIGITS = re.compile(r"\d{4,}\*+\d{2,}|\d{10,}")
_PEOPLE = re.compile(
    r"(?:Gönderen|Sender)\s*:\s*(?P<from>[^,]+?)\s*,\s*(?:Alıcı|Recipient)\s*:\s*(?P<to>[^,]+)",
    re.IGNORECASE,
)


@dataclass(frozen=True)
class Reading:
    label: str  # a short description for the expense
    kind: str  # purchase, bill, card_payment, transfer_out, transfer_in, deposit, refund, other


def _title(name: str) -> str:
    name = " ".join(name.split()).replace("İ", "I")  # "DİLEK".title() keeps a stray dot
    return name.title() if name.isupper() or name.islower() else name


def _without_city(name: str) -> str:
    words = name.split()
    while len(words) > 1 and words[-1].lower() in _CITIES:
        words.pop()
    return " ".join(words)


def understand(t: Transaction) -> Reading:
    """A short label and the kind of line, from the words Turkish banks use (Kuveyt Türk
    and others): "Firma Adı: X,Harcama" is a card purchase at X, "Corporation Name: X,
    Invoice Payment" a bill, "Gönderen / Alıcı" a transfer, "Kredi Kartı Borç Ödemesi" a
    credit-card payment (its purchases are counted already)."""
    text = " ".join(t.description.split())
    lower = text.lower()
    out = t.amount < 0
    if "kredi kartı borç ödemesi" in lower or "kredi karti borc odemesi" in lower:
        return Reading("Credit card payment", "card_payment")
    shop = re.search(r"Firma Adı\s*:\s*(.+?)\s*,\s*Harcama", text, re.IGNORECASE)
    if shop:
        return Reading(_without_city(shop.group(1))[:200], "purchase" if out else "refund")
    bill = re.search(r"Corporation Name\s*:\s*(.+?)\s*,\s*Invoice", text, re.IGNORECASE)
    if bill:
        return Reading(bill.group(1).strip()[:200], "bill" if out else "refund")
    people = _PEOPLE.search(text)
    if people:
        note = text[: people.start()].split(" - ")[0].strip(" -,")
        note = "" if len(note) > 40 or "transfer" in note.lower() else note
        if out:
            label = f"Transfer to {_title(people.group('to').strip())}"
            kind = "transfer_out"
        else:
            label = f"Transfer from {_title(people.group('from').strip())}"
            kind = "transfer_in"
        return Reading(f"{label} ({note})"[:200] if note else label[:200], kind)
    sender = re.match(r"(?:.* - )?(?:Gönderen|Sender)\s*:\s*([^,]+)", text, re.IGNORECASE)
    if sender and not out:
        return Reading(f"Transfer from {_title(sender.group(1).strip())}"[:200], "transfer_in")
    card = re.match(r"0[56],\d{6}\*+\d{4},.*?(?:AM|PM),(?P<shop>[^,]+)", text)
    if card:
        return Reading(card.group("shop").strip()[:200], "purchase" if out else "refund")
    if "para yatırma" in lower or "para yatirma" in lower:
        return Reading("Cash deposit", "deposit" if not out else "other")
    label = " ".join(_LONG_DIGITS.sub("", text).split()).strip(" ,") or "Bank transaction"
    return Reading(label[:200], "other" if out else "transfer_in")
