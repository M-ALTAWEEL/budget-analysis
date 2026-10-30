"""Downloadable reports (feature F5): a month or a year as a PDF or an Excel file, in the
person's language. Rules: docs/BUSINESS_LOGIC.md, section 8.8."""

from __future__ import annotations

import io
from dataclasses import dataclass, field
from datetime import date
from decimal import Decimal
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from reportlab.lib import colors
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle
from reportlab.lib.units import mm
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer, Table, TableStyle

from app.core.arabic import has_arabic, visual

FONTS = Path(__file__).resolve().parents[2] / "assets" / "fonts"
PINE = "#1F5F4A"

_LABELS: dict[str, dict[str, str]] = {
    "en": {
        "title": "Budget report",
        "month": "Month",
        "currency": "Currency",
        "budget": "Budget",
        "spent": "Spent",
        "left": "Left",
        "income": "Income",
        "categories": "Categories",
        "category": "Category",
        "limit": "Limit",
        "expenses": "Expenses",
        "date": "Date",
        "description": "Description",
        "amount": "Amount",
        "original": "Paid as",
        "months": "Months",
        "summary": "Summary",
        "made": "Made with Budget Analysis on",
        "none": "No expenses.",
    },
    "ar": {
        "title": "تقرير الميزانية",
        "month": "الشهر",
        "currency": "العملة",
        "budget": "الميزانية",
        "spent": "المصروف",
        "left": "المتبقي",
        "income": "الدخل",
        "categories": "الفئات",
        "category": "الفئة",
        "limit": "الحد",
        "expenses": "المصاريف",
        "date": "التاريخ",
        "description": "الوصف",
        "amount": "المبلغ",
        "original": "المدفوع",
        "months": "الأشهر",
        "summary": "الملخص",
        "made": "صُنع بواسطة Budget Analysis في",
        "none": "لا توجد مصاريف.",
    },
    "tr": {
        "title": "Bütçe raporu",
        "month": "Ay",
        "currency": "Para birimi",
        "budget": "Bütçe",
        "spent": "Harcanan",
        "left": "Kalan",
        "income": "Gelir",
        "categories": "Kategoriler",
        "category": "Kategori",
        "limit": "Limit",
        "expenses": "Harcamalar",
        "date": "Tarih",
        "description": "Açıklama",
        "amount": "Tutar",
        "original": "Ödenen",
        "months": "Aylar",
        "summary": "Özet",
        "made": "Budget Analysis ile oluşturuldu:",
        "none": "Harcama yok.",
    },
}


def labels(locale: str) -> dict[str, str]:
    return _LABELS.get(locale, _LABELS["en"])


@dataclass
class MonthRow:
    month: str  # "2026-09"
    currency: str
    budget: Decimal
    spent: Decimal
    income: Decimal | None

    @property
    def left(self) -> Decimal:
        return self.budget - self.spent


@dataclass
class CategoryRow:
    name: str
    currency: str
    limit: Decimal | None
    spent: Decimal


@dataclass
class ExpenseRow:
    day: date
    description: str
    category: str
    amount: Decimal  # in the month's currency
    currency: str
    original: str | None  # "10.00 USD" when paid in another currency


@dataclass
class ReportData:
    locale: str
    person: str
    period: str  # "2026-09" or "2026"
    months: list[MonthRow]
    categories: list[CategoryRow]
    expenses: list[ExpenseRow] = field(default_factory=list)


def amount_text(value: Decimal | None, currency: str = "") -> str:
    if value is None:
        return "-"
    return f"{value:,.2f} {currency}".strip()


# ── Excel ────────────────────────────────────────────────────────────────────


def to_xlsx(data: ReportData) -> bytes:
    """Plain numbers in number cells, so the sheet can be summed and charted."""
    words = labels(data.locale)
    book = Workbook()
    rtl = data.locale == "ar"
    head = Font(bold=True, color="FFFFFF")
    fill = PatternFill("solid", fgColor=PINE[1:])

    def sheet(title: str, columns: list[str], rows: list[list], widths: list[int]):
        ws = book.create_sheet(title[:31])
        ws.sheet_view.rightToLeft = rtl
        ws.append(columns)
        for cell in ws[1]:
            cell.font, cell.fill = head, fill
            cell.alignment = Alignment(horizontal="center")
        for row in rows:
            ws.append(row)
        for i, width in enumerate(widths):
            ws.column_dimensions[chr(65 + i)].width = width
        for column in ws.iter_cols(min_row=2):
            for cell in column:
                if isinstance(cell.value, float):
                    cell.number_format = "#,##0.00"
        ws.freeze_panes = "A2"
        return ws

    book.remove(book.active)
    sheet(
        words["summary"],
        [
            words["month"],
            words["currency"],
            words["budget"],
            words["spent"],
            words["left"],
            words["income"],
        ],
        [
            [
                m.month,
                m.currency,
                float(m.budget),
                float(m.spent),
                float(m.left),
                float(m.income) if m.income is not None else None,
            ]
            for m in data.months
        ],
        [12, 10, 14, 14, 14, 14],
    )
    sheet(
        words["categories"],
        [words["category"], words["currency"], words["limit"], words["spent"]],
        [
            [c.name, c.currency, float(c.limit) if c.limit is not None else None, float(c.spent)]
            for c in data.categories
        ],
        [24, 10, 14, 14],
    )
    sheet(
        words["expenses"],
        [
            words["date"],
            words["description"],
            words["category"],
            words["amount"],
            words["currency"],
            words["original"],
        ],
        [
            [e.day, e.description, e.category, float(e.amount), e.currency, e.original]
            for e in data.expenses
        ],
        [12, 36, 18, 14, 10, 16],
    )
    out = io.BytesIO()
    book.save(out)
    return out.getvalue()


# ── PDF ──────────────────────────────────────────────────────────────────────

_registered = False


def _fonts() -> None:
    global _registered
    if not _registered:
        pdfmetrics.registerFont(TTFont("Report", str(FONTS / "DejaVuSans.ttf")))
        pdfmetrics.registerFont(TTFont("Report-Bold", str(FONTS / "DejaVuSans-Bold.ttf")))
        _registered = True


def to_pdf(data: ReportData) -> bytes:
    _fonts()
    words = labels(data.locale)
    rtl = data.locale == "ar"
    align = 2 if rtl else 0  # TA_RIGHT / TA_LEFT

    def para(value: str, size: float = 9, bold: bool = False, color=colors.black, al=None):
        style = ParagraphStyle(
            "p",
            fontName="Report-Bold" if bold else "Report",
            fontSize=size,
            leading=size * 1.35,
            textColor=color,
            alignment=(2 if has_arabic(value) else 0) if al is None else al,
        )
        safe = visual(value).replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")
        return Paragraph(safe, style)

    def table(columns: list[str], rows: list[list[str]], widths: list[float], numeric: set[int]):
        if rtl:  # columns read from the right
            columns, rows = columns[::-1], [r[::-1] for r in rows]
            widths = widths[::-1]
            numeric = {len(widths) - 1 - i for i in numeric}

        def cell_align(i: int) -> int:
            return 2 if rtl or i in numeric else 0

        body = [
            [
                para(c, 8.5, bold=True, color=colors.white, al=cell_align(i))
                for i, c in enumerate(columns)
            ]
        ]
        body += [[para(v, 8.5, al=cell_align(i)) for i, v in enumerate(row)] for row in rows]
        t = Table(body, colWidths=widths, repeatRows=1)
        t.setStyle(
            TableStyle(
                [
                    ("BACKGROUND", (0, 0), (-1, 0), colors.HexColor(PINE)),
                    (
                        "ROWBACKGROUNDS",
                        (0, 1),
                        (-1, -1),
                        [colors.white, colors.HexColor("#F3F5F2")],
                    ),
                    ("VALIGN", (0, 0), (-1, -1), "MIDDLE"),
                    ("TOPPADDING", (0, 0), (-1, -1), 4),
                    ("BOTTOMPADDING", (0, 0), (-1, -1), 4),
                    ("LINEBELOW", (0, 0), (-1, -1), 0.25, colors.HexColor("#DADFD8")),
                ]
            )
        )
        return t

    out = io.BytesIO()
    doc = SimpleDocTemplate(
        out,
        pagesize=A4,
        leftMargin=16 * mm,
        rightMargin=16 * mm,
        topMargin=16 * mm,
        bottomMargin=16 * mm,
        title=f"{words['title']} {data.period}",
        author="Budget Analysis",
    )
    width = A4[0] - 32 * mm
    story: list = [
        para(
            f"{words['title']} · {data.period}",
            18,
            bold=True,
            color=colors.HexColor(PINE),
            al=align,
        ),
        Spacer(1, 2 * mm),
        para(data.person, 10, color=colors.HexColor("#5A6360"), al=align),
        Spacer(1, 6 * mm),
        para(words["summary"], 12, bold=True, al=align),
        Spacer(1, 2 * mm),
        table(
            [words["month"], words["budget"], words["spent"], words["left"], words["income"]],
            [
                [
                    m.month,
                    amount_text(m.budget, m.currency),
                    amount_text(m.spent, m.currency),
                    amount_text(m.left, m.currency),
                    amount_text(m.income, m.currency),
                ]
                for m in data.months
            ],
            [width * 0.16, width * 0.21, width * 0.21, width * 0.21, width * 0.21],
            {1, 2, 3, 4},
        ),
        Spacer(1, 6 * mm),
        para(words["categories"], 12, bold=True, al=align),
        Spacer(1, 2 * mm),
        table(
            [words["category"], words["limit"], words["spent"]],
            [
                [c.name, amount_text(c.limit, c.currency), amount_text(c.spent, c.currency)]
                for c in data.categories
            ],
            [width * 0.5, width * 0.25, width * 0.25],
            {1, 2},
        ),
        Spacer(1, 6 * mm),
        para(words["expenses"], 12, bold=True, al=align),
        Spacer(1, 2 * mm),
    ]
    if data.expenses:
        story.append(
            table(
                [words["date"], words["description"], words["category"], words["amount"]],
                [
                    [
                        e.day.isoformat(),
                        e.description,
                        e.category,
                        amount_text(e.amount, e.currency)
                        + (f" ({e.original})" if e.original else ""),
                    ]
                    for e in data.expenses
                ],
                [width * 0.15, width * 0.4, width * 0.18, width * 0.27],
                {3},
            )
        )
    else:
        story.append(para(words["none"], 9, al=align))
    story += [
        Spacer(1, 8 * mm),
        para(f"{words['made']} {date.today().isoformat()}", 7.5, color=colors.grey, al=align),
    ]
    doc.build(story)
    return out.getvalue()
