"""Feature F5: a month or a year as a PDF or Excel file."""

import io

from openpyxl import load_workbook

from app.core.arabic import shape, visual
from tests.helpers import API, category_uuid, set_budget, signup


async def a_september(client, **signup_args) -> dict:
    me = await signup(client, **signup_args)
    await set_budget(client, me, "2026-09", total="4000")
    await set_budget(client, me, "2026-08", total="5000")
    food = await category_uuid(client, me, "food")
    await client.put(
        f"{API}/budgets/2026-09/allocations",
        headers=me,
        json={"allocations": [{"category_uuid": food, "limit_amount": "1200"}]},
    )
    for amount, currency, day in (("300", None, "2026-09-01"), ("10", "USD", "2026-09-05")):
        body = {"amount": amount, "description": "Market", "category_uuid": food, "spent_at": day}
        if currency:
            body["currency"] = currency
        response = await client.post(f"{API}/expenses", headers=me, json=body)
        assert response.status_code == 201, response.text
    return me


async def test_the_month_as_an_excel_file(client):
    me = await a_september(client)
    response = await client.get(f"{API}/reports/2026-09/file?format=xlsx", headers=me)
    assert response.status_code == 200
    assert response.headers["content-disposition"] == 'attachment; filename="budget-2026-09.xlsx"'
    assert response.headers["cache-control"] == "no-store"
    book = load_workbook(io.BytesIO(response.content))
    assert book.sheetnames == ["Summary", "Categories", "Expenses"]
    summary = list(book["Summary"].iter_rows(min_row=2, values_only=True))
    assert summary == [("2026-09", "TRY", 4000, 640, 3360, None)]
    categories = list(book["Categories"].iter_rows(min_row=2, values_only=True))
    assert categories == [("Food", "TRY", 1200, 640)]
    expenses = list(book["Expenses"].iter_rows(min_row=2, values_only=True))
    assert [(e[1], e[3], e[5]) for e in expenses] == [
        ("Market", 300, None),
        ("Market", 340, "10.00 USD"),
    ]


async def test_the_year_as_a_pdf_in_arabic(client):
    me = await a_september(client)
    await client.patch(f"{API}/users/me", headers=me, json={"locale": "ar"})
    response = await client.get(f"{API}/reports/year/2026/file?format=pdf", headers=me)
    assert response.status_code == 200, response.text
    assert response.headers["content-type"] == "application/pdf"
    assert response.content.startswith(b"%PDF")


async def test_no_budget_no_file_and_only_known_formats(client):
    me = await signup(client)
    assert (
        await client.get(f"{API}/reports/2026-09/file?format=pdf", headers=me)
    ).status_code == 404
    assert (
        await client.get(f"{API}/reports/year/2026/file?format=pdf", headers=me)
    ).status_code == 404
    assert (
        await client.get(f"{API}/reports/2026-09/file?format=doc", headers=me)
    ).status_code == 422
    assert (await client.get(f"{API}/reports/2026-09/file?format=pdf")).status_code == 401


def test_arabic_is_joined_and_reads_right_to_left():
    assert shape("طعام") == "ﻃﻌﺎﻡ"  # initial, medial, final, alone
    assert shape("لا") == "ﻻ"  # lam-alef ligature
    assert visual("Budget") == "Budget"
    assert visual("الميزانية 1,250.00 TRY").startswith("1,250.00 TRY ")
